"""模块四：智能分析报告生成与导出。

文字由模板骨架生成，所有数字从分析结果对象注入，并保留来源标记；
可选的 LLM 只做措辞润色，不改动数值。
"""

from __future__ import annotations

import os
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Callable

from .analytics import overview
from .config import Settings
from .models import AnalysisResult, ChartSpec, CleaningLog, IngestReport
from .plans import section_chart_map


@dataclass
class ReportSection:
    id: str
    title: str
    paragraphs: list[str] = field(default_factory=list)
    tables: list[dict[str, Any]] = field(default_factory=list)
    charts: list[ChartSpec] = field(default_factory=list)


@dataclass
class ReportBundle:
    title: str
    subtitle: str
    generated_at: str
    meta: dict[str, Any] = field(default_factory=dict)
    sections: list[ReportSection] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "title": self.title,
            "subtitle": self.subtitle,
            "generated_at": self.generated_at,
            "meta": self.meta,
            "sections": [
                {
                    "id": section.id,
                    "title": section.title,
                    "paragraphs": section.paragraphs,
                    "tables": section.tables,
                    "charts": [chart.to_dict() for chart in section.charts],
                }
                for section in self.sections
            ],
        }


def _fmt(value: Any) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value)
    if abs(number) >= 10000:
        return f"{number:,.0f}"
    if number == int(number):
        return f"{int(number):,}"
    return f"{number:,.2f}"


def _top_records(records: list[dict[str, Any]], key: str, count: int = 3) -> list[dict[str, Any]]:
    return sorted(records, key=lambda item: item.get(key) or 0, reverse=True)[:count]


def _metric_total(result: AnalysisResult | None, label: str, rows: list[dict[str, Any]]) -> float:
    """取指标总量。

    分析结果里的 records 可能被 limit 截断，直接求和会把分母算小，
    因此优先用未被截断的 totals。
    """
    if result is not None:
        value = result.totals.get(label)
        if value is not None:
            return float(value)
    return sum((row.get(label) or 0) for row in rows) or 1e-9


class ReportBuilder:
    """把分析结果组织成报告结构。"""

    def __init__(
        self,
        settings: Settings,
        frame: Any,
        ingest_report: IngestReport,
        cleaning_log: CleaningLog,
        analyses: dict[str, AnalysisResult],
        charts: dict[str, ChartSpec],
        polish: Callable[[str, str], str] | None = None,
    ) -> None:
        self.settings = settings
        self.frame = frame
        self.ingest_report = ingest_report
        self.cleaning_log = cleaning_log
        self.analyses = analyses
        self.charts = charts
        self.polish = polish
        self.chart_sections = section_chart_map()

    # ---- 生成 ----
    def build(self) -> ReportBundle:
        bundle = ReportBundle(
            title=self.settings.outline["title"],
            subtitle=self.settings.outline["subtitle"],
            generated_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            meta={
                "数据文件": "、".join(Path(name).name for name in self.ingest_report.files),
                "原始记录数": self.ingest_report.rows_read,
                "清洗后记录数": self.cleaning_log.rows_out,
                "清洗动作数": len(self.cleaning_log.steps),
                "待确认问题数": len(self.cleaning_log.issues),
            },
        )
        for section_spec in self.settings.outline["sections"]:
            generator = getattr(self, f"_section_{section_spec['generator']}", None)
            if generator is None:
                continue
            section = ReportSection(id=section_spec["id"], title=section_spec["title"])
            generator(section)
            for key in self.chart_sections.get(section_spec["id"], []):
                chart = self.charts.get(key)
                if chart is not None:
                    section.charts.append(chart)
            if self.polish:
                section.paragraphs = [self.polish(section.title, text) for text in section.paragraphs]
            bundle.sections.append(section)
        bundle.sections.append(self._section_data_quality())
        return bundle

    # ---- 各章节 ----
    def _section_overview(self, section: ReportSection) -> None:
        stats = overview(self.frame, self.settings)
        section.paragraphs.append(
            "本报告基于 {files} 的灾情数据自动生成，统计时段为 {period}，"
            "清洗后共纳入有效记录 {rows} 条，覆盖 {types} 个灾种、{counties} 个县（区）。"
            "报告由数据接入、清洗标准化、多维分析、可视化与报告生成五个模块自动完成，"
            "全部数值来自分析结果对象，未经人工修改。".format(
                files="、".join(Path(name).name for name in self.ingest_report.files),
                period=stats.get("统计时段", "未知"),
                rows=_fmt(stats.get("记录条数", 0)),
                types=_fmt(stats.get("涉及灾种数", 0)),
                counties=_fmt(stats.get("覆盖县（区）数", 0)),
            )
        )
        section.paragraphs.append(
            "总体来看，统计期内受灾人口 {affected} 人，因灾死亡失踪 {death} 人，"
            "紧急转移安置 {evacuated} 人，直接经济损失合计 {loss} 万元，"
            "单次事件平均直接经济损失 {average} 万元。".format(
                affected=_fmt(stats.get("受灾人口（人）", 0)),
                death=_fmt(stats.get("因灾死亡失踪人口（人）", 0)),
                evacuated=_fmt(stats.get("紧急转移安置人口（人）", 0)),
                loss=_fmt(stats.get("直接经济损失（万元）", 0)),
                average=_fmt(stats.get("单次事件平均直接经济损失（万元）", 0)),
            )
        )
        section.tables.append(
            {
                "caption": "表 1 核心指标汇总",
                "columns": ["指标", "数值"],
                "rows": [[key, _fmt(value)] for key, value in stats.items()],
            }
        )

    def _section_spatiotemporal(self, section: ReportSection) -> None:
        trend = self.analyses.get("yearly_trend")
        county = self.analyses.get("county_ranking")
        by_type_trend = self.analyses.get("yearly_trend_by_type")

        if trend and trend.records:
            label = "年份"
            values = [item for item in trend.records if item.get(label) is not None]
            peak = max(values, key=lambda item: item.get("直接经济损失") or 0) if values else None
            if peak:
                section.paragraphs.append(
                    "从时间分布看，数据覆盖 {start} 年至 {end} 年。损失最重的年份为 {year} 年，"
                    "直接经济损失 {value} 万元，占统计期总量的 {share}%。".format(
                        start=int(values[0][label]),
                        end=int(values[-1][label]),
                        year=int(peak[label]),
                        value=_fmt(peak.get("直接经济损失", 0)),
                        share=round(
                            (peak.get("直接经济损失") or 0)
                            / max(sum(item.get("直接经济损失") or 0 for item in values), 1e-9)
                            * 100,
                            2,
                        ),
                    )
                )
            if len(values) >= 3:
                recent = values[-1].get("直接经济损失") or 0
                previous = values[-2].get("直接经济损失") or 0
                direction = "上升" if recent > previous else "下降"
                section.paragraphs.append(
                    f"最近一个统计年度损失较上年{direction}，由 {_fmt(previous)} 万元变为 {_fmt(recent)} 万元。"
                )

        if by_type_trend and by_type_trend.records:
            months = self.analyses.get("by_season")
            if months and months.records:
                flood = next((item for item in months.records if item.get("汛期标记") == "汛期"), None)
                total = sum(item.get("直接经济损失") or 0 for item in months.records) or 1e-9
                if flood:
                    section.paragraphs.append(
                        "汛期（5—9 月）直接经济损失 {value} 万元，占全年总量的 {share}%，"
                        "灾情发生具有明显的汛期集中特征。".format(
                            value=_fmt(flood.get("直接经济损失", 0)),
                            share=round((flood.get("直接经济损失") or 0) / total * 100, 2),
                        )
                    )

        if county and county.records:
            top = _top_records(county.records, "直接经济损失", 3)
            total = _metric_total(county, "直接经济损失", county.records)
            detail = "、".join(
                f"{item.get('县（区）')}（{_fmt(item.get('直接经济损失', 0))} 万元）" for item in top
            )
            section.paragraphs.append(
                "从区域分布看，损失最集中的三个县（区）为{detail}，"
                "合计占统计期总损失的 {share}%。".format(
                    detail=detail,
                    share=round(sum(item.get("直接经济损失") or 0 for item in top) / total * 100, 2),
                )
            )
        if not section.paragraphs:
            section.paragraphs.append("当前数据不足以形成时空分布结论，建议补充时间或区域字段后重跑。")

    def _section_by_disaster_type(self, section: ReportSection) -> None:
        by_type = self.analyses.get("by_disaster_type")
        if not by_type or not by_type.records:
            section.paragraphs.append("当前数据不足以形成分灾种结论。")
            return
        rows = by_type.records
        total = _metric_total(by_type, "直接经济损失", rows)
        leader = max(rows, key=lambda item: item.get("直接经济损失") or 0)
        section.paragraphs.append(
            "统计期内共识别 {count} 类灾种，直接经济损失共 {total} 万元。"
            "其中 {leader} 造成的损失最大，为 {value} 万元，占比 {share}%。".format(
                count=len(rows),
                total=_fmt(total),
                leader=leader.get("灾种"),
                value=_fmt(leader.get("直接经济损失", 0)),
                share=round((leader.get("直接经济损失") or 0) / total * 100, 2),
            )
        )
        second = sorted(rows, key=lambda item: item.get("直接经济损失") or 0, reverse=True)[1:3]
        if second:
            section.paragraphs.append(
                "其余主要灾种为"
                + "、".join(
                    f"{item.get('灾种')}（{_fmt(item.get('直接经济损失', 0))} 万元）" for item in second
                )
                + "，灾种结构呈现多灾种并发的特征。"
            )
        section.tables.append(
            {
                "caption": "表 2 分灾种直接经济损失",
                "columns": ["灾种", "直接经济损失（万元）", "占比"],
                "rows": [
                    [
                        item.get("灾种"),
                        _fmt(item.get("直接经济损失", 0)),
                        f"{round((item.get('直接经济损失') or 0) / total * 100, 2)}%",
                    ]
                    for item in rows
                ],
            }
        )

    def _section_loss_structure(self, section: ReportSection) -> None:
        structure = self.analyses.get("loss_structure")
        if structure and structure.records:
            amount_rows = [item for item in structure.records if item.get("单位") == "万元"]
            if amount_rows:
                leader = max(amount_rows, key=lambda item: item.get("数值") or 0)
                section.paragraphs.append(
                    "损失结构方面，直接经济损失合计 {total} 万元。其中 {name} 占比最高，"
                    "为 {value} 万元，占 {share}%，反映出该类损失是灾后恢复重建的主要压力来源。".format(
                        total=_fmt(structure.totals.get("损失合计（万元）", 0)),
                        name=leader.get("损失类型"),
                        value=_fmt(leader.get("数值", 0)),
                        share=round((leader.get("占比") or 0) * 100, 2),
                    )
                )
            people = next(
                (item for item in structure.records if item.get("损失类型") == "人员伤亡"), None
            )
            if people:
                section.paragraphs.append(
                    f"人员伤亡方面，因灾死亡失踪与受伤人口合计 {_fmt(people.get('数值', 0))} 人，"
                    "需重点关注人员转移安置与生命线工程的保障能力。"
                )
            section.tables.append(
                {
                    "caption": "表 3 损失结构构成",
                    "columns": ["损失类型", "数值", "单位", "占比"],
                    "rows": [
                        [
                            item.get("损失类型"),
                            _fmt(item.get("数值", 0)),
                            item.get("单位"),
                            f"{round((item.get('占比') or 0) * 100, 2)}%"
                            if item.get("占比") is not None
                            else "—",
                        ]
                        for item in structure.records
                    ],
                }
            )
        season = self.analyses.get("by_season")
        if season and season.records:
            for item in season.records:
                section.paragraphs.append(
                    "{season}受灾人口 {affected} 人，直接经济损失 {loss} 万元。".format(
                        season=item.get("汛期标记"),
                        affected=_fmt(item.get("受灾人口", 0)),
                        loss=_fmt(item.get("直接经济损失", 0)),
                    )
                )

    def _section_conclusion(self, section: ReportSection) -> None:
        by_type = self.analyses.get("by_disaster_type")
        county = self.analyses.get("county_ranking")
        structure = self.analyses.get("loss_structure")
        season = self.analyses.get("by_season")

        conclusions: list[str] = []
        if by_type and by_type.records:
            leader = max(by_type.records, key=lambda item: item.get("直接经济损失") or 0)
            by_type_total = _metric_total(by_type, "直接经济损失", by_type.records)
            conclusions.append(
                f"主导灾种为{leader.get('灾种')}，其直接经济损失占统计期总量的 "
                f"{round((leader.get('直接经济损失') or 0) / by_type_total * 100, 2)}%。"
            )
        if county and county.records:
            top = _top_records(county.records, "直接经济损失", 3)
            conclusions.append(
                "损失重心集中在"
                + "、".join(str(item.get("县（区）")) for item in top)
                + "，建议将上述区域列为重点防灾减灾单元。"
            )
        if season and season.records:
            flood = next((item for item in season.records if item.get("汛期标记") == "汛期"), None)
            if flood:
                total = sum(item.get("直接经济损失") or 0 for item in season.records) or 1e-9
                conclusions.append(
                    f"汛期损失占比达 {round((flood.get('直接经济损失') or 0) / total * 100, 2)}%，"
                    "防汛关键期应提前部署力量前置与物资预置。"
                )
        if structure and structure.records:
            amount_rows = [item for item in structure.records if item.get("单位") == "万元"]
            if amount_rows:
                leader = max(amount_rows, key=lambda item: item.get("数值") or 0)
                conclusions.append(
                    f"损失结构中{leader.get('损失类型')}占比最高（{round((leader.get('占比') or 0) * 100, 2)}%），"
                    "恢复重建资金应向该方向倾斜。"
                )

        cleaned = [item.rstrip("。；") for item in conclusions]
        section.paragraphs.append("主要结论：" + "；".join(cleaned) + "。")

        suggestions = [
            "（一）预案编制：结合分灾种、分区域、分时段分析结果，细化重点县（区）与汛期专项预案，明确响应启动条件与处置流程。",
            "（二）应急响应：针对主导灾种与高发时段，前置救援力量与应急物资；发生灾情后按损失结构优先级安排抢通保通与人员安置。",
            "（三）风险防控：将损失最集中的县（区）纳入重点监测名录，加强隐患点巡查与预警信息发布；对重复受灾区域开展工程治理评估。",
            "（四）数据治理：对报告中标记为待确认的行政区划与缺失字段逐条核实，提升灾情报送的完整性与一致性，为后续分析提供可靠底数。",
        ]
        section.paragraphs.append("应急管理工作建议：")
        section.paragraphs.extend(suggestions)

    def _section_data_quality(self) -> ReportSection:
        section = ReportSection(id="data_quality", title="附录：数据质量与处理说明")
        ingest = self.ingest_report
        section.paragraphs.append(
            "字段比对：共读取 {sheets} 个工作表，原始记录 {rows_in} 条，"
            "成功映射标准字段 {matched} 个，未识别列 {unmapped} 个，字段匹配率 {rate}%。".format(
                sheets=len(ingest.sheets),
                rows_in=ingest.rows_read,
                matched=len(ingest.matched_columns),
                unmapped=len(ingest.unmapped_columns),
                rate=round(ingest.matched_rate * 100, 2),
            )
        )
        if ingest.unmapped_columns:
            section.paragraphs.append(
                "未识别列（需确认是否补充字段映射）：" + "、".join(ingest.unmapped_columns[:12])
            )
        if ingest.missing_required:
            section.paragraphs.append(
                "缺失的必需字段：" + "、".join(ingest.missing_required) + "，建议在模板中补齐后重新导入。"
            )
        section.paragraphs.append(
            "清洗动作共 {steps} 项，记录数由 {before} 条变为 {after} 条。".format(
                steps=len(self.cleaning_log.steps),
                before=self.cleaning_log.rows_in,
                after=self.cleaning_log.rows_out,
            )
        )
        section.tables.append(
            {
                "caption": "表 4 清洗日志",
                "columns": ["环节", "字段", "处理方式", "影响记录数"],
                "rows": [
                    [step.step, step.field, step.action, step.affected]
                    for step in self.cleaning_log.steps
                ],
            }
        )
        if self.cleaning_log.issues:
            section.tables.append(
                {
                    "caption": "表 5 待确认数据质量问题",
                    "columns": ["级别", "问题", "记录数", "示例"],
                    "rows": [
                        [
                            issue.level,
                            issue.message,
                            issue.count,
                            "；".join(issue.samples[:3]),
                        ]
                        for issue in self.cleaning_log.issues
                    ],
                }
            )
        return section


# ---------------- 导出 ----------------


def write_markdown(bundle: ReportBundle, path: str | Path) -> Path:
    path = Path(path)
    lines = [f"# {bundle.title}", f"_{bundle.subtitle}_", "", f"生成时间：{bundle.generated_at}", ""]
    lines.append("| 项目 | 内容 |")
    lines.append("| --- | --- |")
    for key, value in bundle.meta.items():
        lines.append(f"| {key} | {value} |")
    for section in bundle.sections:
        lines.extend(["", f"## {section.title}", ""])
        for paragraph in section.paragraphs:
            lines.extend([paragraph, ""])
        for table in section.tables:
            lines.append(f"**{table['caption']}**")
            lines.append("")
            lines.append("| " + " | ".join(str(item) for item in table["columns"]) + " |")
            lines.append("| " + " | ".join("---" for _ in table["columns"]) + " |")
            for row in table["rows"]:
                lines.append("| " + " | ".join(str(item) for item in row) + " |")
            lines.append("")
        for chart in section.charts:
            if chart.png_path:
                relative = os.path.relpath(chart.png_path, path.parent).replace("\\", "/")
                lines.extend([f"![{chart.title}]({relative})", "", f"*{chart.label}：{chart.why}*", ""])
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def write_docx(bundle: ReportBundle, path: str | Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    document = Document()
    heading = document.add_heading(bundle.title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(bundle.subtitle)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(f"生成时间：{bundle.generated_at}")
    for key, value in bundle.meta.items():
        document.add_paragraph(f"{key}：{value}", style="List Bullet")

    for section in bundle.sections:
        document.add_heading(section.title, level=1)
        for paragraph in section.paragraphs:
            text = document.add_paragraph(paragraph)
            text.paragraph_format.first_line_indent = Pt(21)
            text.paragraph_format.space_after = Pt(6)
        for table in section.tables:
            caption = document.add_paragraph(table["caption"])
            caption.runs[0].bold = True
            rows = table["rows"][:60]
            grid = document.add_table(rows=1, cols=len(table["columns"]))
            grid.style = "Table Grid"
            for index, name in enumerate(table["columns"]):
                cell = grid.rows[0].cells[index]
                cell.text = str(name)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
            for row in rows:
                cells = grid.add_row().cells
                for index, value in enumerate(row):
                    cells[index].text = str(value)
        for chart in section.charts:
            if chart.png_path and Path(chart.png_path).exists():
                document.add_picture(chart.png_path, width=Cm(15.5))
                note = document.add_paragraph(f"{chart.title}（{chart.label}）")
                note.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.save(str(path))
    return Path(path)


def write_pdf(bundle: ReportBundle, path: str | Path) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from reportlab.lib import colors

    font_name = "Helvetica"
    try:
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        font_name = "STSong-Light"
    except Exception:
        pass

    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontName=font_name, fontSize=10.5, leading=17)
    title_style = ParagraphStyle(
        "title", parent=styles["Title"], fontName=font_name, fontSize=20, leading=28
    )
    heading_style = ParagraphStyle(
        "heading", parent=styles["Heading1"], fontName=font_name, fontSize=14, leading=22, spaceBefore=14
    )
    caption_style = ParagraphStyle(
        "caption", parent=styles["BodyText"], fontName=font_name, fontSize=10, leading=15
    )

    document = SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=2.0 * cm, bottomMargin=2.0 * cm
    )
    flow: list[Any] = [
        Paragraph(bundle.title, title_style),
        Paragraph(bundle.subtitle, ParagraphStyle("sub", parent=body, alignment=1)),
        Spacer(1, 10),
        Paragraph(f"生成时间：{bundle.generated_at}", body),
        Spacer(1, 6),
    ]
    for key, value in bundle.meta.items():
        flow.append(Paragraph(f"· {key}：{value}", body))

    for section in bundle.sections:
        flow.append(Paragraph(section.title, heading_style))
        for paragraph in section.paragraphs:
            flow.append(Paragraph(str(paragraph), body))
        for table in section.tables:
            flow.append(Paragraph(f"<b>{table['caption']}</b>", caption_style))
            data = [list(table["columns"])] + [
                [str(item) for item in row] for row in table["rows"][:40]
            ]
            grid = Table(data, repeatRows=1, hAlign="LEFT")
            grid.setStyle(
                TableStyle(
                    [
                        ("FONTNAME", (0, 0), (-1, -1), font_name),
                        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9aa5b1")),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef6")),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 4),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            flow.extend([grid, Spacer(1, 8)])
        for chart in section.charts:
            if chart.png_path and Path(chart.png_path).exists():
                max_width = 15.5 * cm
                max_height = 16.0 * cm
                width = max_width
                height = width * 0.62
                try:
                    from PIL import Image as PILImage

                    with PILImage.open(chart.png_path) as handle:
                        raw_width, raw_height = handle.size
                    ratio = raw_height / max(raw_width, 1)
                    height = width * ratio
                    if height > max_height:
                        height = max_height
                        width = height / ratio
                except Exception:
                    pass
                flow.append(Image(chart.png_path, width=width, height=height))
                flow.append(Paragraph(f"{chart.title}（{chart.label}）", caption_style))
                flow.append(Spacer(1, 8))

    document.build(flow)
    return Path(path)
